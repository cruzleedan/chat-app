"""Document processing and chunking logic."""

from pathlib import Path
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_community.document_loaders import (
    PyPDFLoader, 
    TextLoader, 
    Docx2txtLoader,
    BSHTMLLoader,
)
from langchain_core.documents import Document

from app.config import settings


class DocumentProcessor:
    """Handles document loading and chunking."""

    SUPPORTED_EXTENSIONS = {
        ".pdf", ".txt", ".md", ".docx", ".doc", ".html", ".htm",
        ".xlsx", ".xls",  # Excel
        ".pptx", ".ppt",  # PowerPoint
        ".jpg", ".jpeg", ".png",  # Images (OCR)
    }

    def __init__(self):
        """Initialize the document processor with text splitter."""
        self.text_splitter = RecursiveCharacterTextSplitter(
            chunk_size=settings.chunk_size,
            chunk_overlap=settings.chunk_overlap,
            length_function=len,
            separators=["\n\n", "\n", ". ", " ", ""],
        )

    def is_supported(self, filename: str) -> bool:
        """Check if file type is supported."""
        ext = Path(filename).suffix.lower()
        return ext in self.SUPPORTED_EXTENSIONS

    def _load_excel(self, file_path: Path) -> list[Document]:
        """Load an Excel file using pandas."""
        import pandas as pd
        
        ext = file_path.suffix.lower()
        documents = []
        
        # Read all sheets
        if ext == ".xlsx":
            excel_file = pd.ExcelFile(file_path, engine="openpyxl")
        else:  # .xls
            excel_file = pd.ExcelFile(file_path, engine="xlrd")
        
        for sheet_name in excel_file.sheet_names:
            df = pd.read_excel(excel_file, sheet_name=sheet_name)
            
            # Convert DataFrame to readable text
            content_parts = [f"Sheet: {sheet_name}\n"]
            content_parts.append(df.to_string(index=False))
            
            documents.append(Document(
                page_content="\n".join(content_parts),
                metadata={"sheet": sheet_name}
            ))
        
        return documents

    def _load_powerpoint(self, file_path: Path) -> list[Document]:
        """Load a PowerPoint file using python-pptx."""
        from pptx import Presentation
        
        prs = Presentation(str(file_path))
        documents = []
        
        for slide_num, slide in enumerate(prs.slides, start=1):
            text_parts = []
            
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    text_parts.append(shape.text)
            
            if text_parts:
                documents.append(Document(
                    page_content="\n".join(text_parts),
                    metadata={"page": slide_num}
                ))
        
        return documents

    def _load_image(self, file_path: Path) -> list[Document]:
        """Load an image file using OCR (pytesseract)."""
        try:
            import pytesseract
            from PIL import Image
            
            image = Image.open(file_path)
            text = pytesseract.image_to_string(image)
            
            if not text.strip():
                text = "[Image contains no extractable text - the image may contain graphics, diagrams, or non-text content]"
            else:
                # Prepend context so LLM knows this text was extracted from an image
                filename = file_path.name
                text = f"[The following text was extracted from an image file named '{filename}' using OCR (Optical Character Recognition). This is the visible text content from that image:]\n\n{text}"
            
            return [Document(
                page_content=text,
                metadata={"source": str(file_path), "content_type": "image_ocr"}
            )]
        except Exception as e:
            # If OCR fails, return a placeholder
            return [Document(
                page_content=f"[Image file - OCR extraction failed: {str(e)}]",
                metadata={"source": str(file_path), "content_type": "image_ocr"}
            )]

    def load_document(self, file_path: Path) -> list[Document]:
        """
        Load a document from file path.

        Args:
            file_path: Path to the document file

        Returns:
            List of Document objects

        Raises:
            ValueError: If file type is not supported
        """
        ext = file_path.suffix.lower()

        if ext == ".pdf":
            loader = PyPDFLoader(str(file_path))
        elif ext in {".txt", ".md"}:
            loader = TextLoader(str(file_path))
        elif ext in {".docx", ".doc"}:
            loader = Docx2txtLoader(str(file_path))
        elif ext in {".html", ".htm"}:
            loader = BSHTMLLoader(str(file_path))
        elif ext in {".xlsx", ".xls"}:
            return self._load_excel(file_path)
        elif ext in {".pptx", ".ppt"}:
            return self._load_powerpoint(file_path)
        elif ext in {".jpg", ".jpeg", ".png"}:
            return self._load_image(file_path)
        else:
            raise ValueError(f"Unsupported file type: {ext}")

        return loader.load()

    def chunk_documents(
        self, documents: list[Document], document_id: str, filename: str
    ) -> list[Document]:
        """
        Split documents into chunks and add metadata.

        Args:
            documents: List of loaded documents
            document_id: Unique identifier for the source document
            filename: Original filename

        Returns:
            List of chunked Document objects with metadata
        """
        chunks = self.text_splitter.split_documents(documents)

        # Add metadata to each chunk
        for i, chunk in enumerate(chunks):
            chunk.metadata.update({
                "document_id": document_id,
                "filename": filename,
                "chunk_index": i,
                "chunk_id": f"{document_id}_chunk_{i}",
            })

            # Ensure page number is present (default to 0 for non-PDF)
            if "page" not in chunk.metadata:
                chunk.metadata["page"] = 0

        return chunks

    def process_file(
        self, file_path: Path, document_id: str, original_filename: str = None
    ) -> list[Document]:
        """
        Full pipeline: load and chunk a document.

        Args:
            file_path: Path to the document
            document_id: Unique identifier for the document
            original_filename: Original filename (if different from file_path.name)

        Returns:
            List of processed chunks
        """
        documents = self.load_document(file_path)
        filename = original_filename or file_path.name
        chunks = self.chunk_documents(
            documents, document_id, filename
        )
        return chunks


# Global instance
document_processor = DocumentProcessor()
